"""Create an Excel report for cooperative rotation measurements."""

from dataclasses import dataclass, field
from datetime import datetime
import math
from pathlib import Path
from statistics import fmean
from tempfile import NamedTemporaryFile
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


FORCE_POINTS = ('amir1_left', 'amir1_right', 'amir2_left', 'amir2_right')
RIGHT_FINGER_POINTS = ('amir1_right', 'amir2_right')
FORCE_LABELS = {
    'amir1_left': 'AMIR1 左指',
    'amir1_right': 'AMIR1 右指',
    'amir2_left': 'AMIR2 左指',
    'amir2_right': 'AMIR2 右指',
}
DISPLAY_FORCE_FILTER_ALPHA = 0.1
CALIBRATED_NORMAL_FORCE_COLUMNS = (20, 30, 40, 50)
SMOOTHED_NORMAL_FORCE_COLUMNS = (51, 52, 53, 54)
PAYLOAD_CENTER_X_COLUMN = 55
PAYLOAD_CENTER_Y_COLUMN = 56
PAYLOAD_ROLL_COLUMN = 57
PAYLOAD_PITCH_COLUMN = 58
PAYLOAD_YAW_COLUMN = 59
SEQUENCE_COLORS = {
    'ロボット・物体生成': '#D9EAF7',
    '把持': '#FCE4D6',
    '把持完了・搬送準備': '#E2F0D9',
    '搬送（直進）': '#DDEBF7',
    '搬送（回転）': '#E4DFEC',
    '搬送（Y方向直進）': '#FFF2CC',
    '搬送完了': '#FFF2CC',
    '異常終了': '#F4CCCC',
}
PLOT_COLORS = ('#1F4E78', '#C55A11', '#548235', '#7030A0')
RIGHT_FINGER_COLORS = {
    'amir1_right': '#C55A11',
    'amir2_right': '#7030A0',
}
RIGHT_FINGER_TIME_TICK_SECONDS = 50
FONT_PATH = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'


def _use_five_unit_axis_labels(
        chart: ScatterChart, final_elapsed: float) -> None:
    """Use true five-unit numeric axes without rounding source data."""
    chart.x_axis.numFmt = '0'
    chart.x_axis.majorUnit = 5
    chart.x_axis.scaling.min = 0
    chart.x_axis.scaling.max = max(5, math.ceil(final_elapsed / 5.0) * 5)
    chart.y_axis.numFmt = '0'
    chart.y_axis.majorUnit = 5


def _add_xy_series(
        chart: ScatterChart, data_sheet, columns: Tuple[int, ...],
        times: Reference, last_row: int) -> None:
    """Add one numeric-X series per requested data column."""
    for column in columns:
        values = Reference(
            data_sheet, min_col=column, min_row=1, max_row=last_row)
        chart.series.append(Series(values, xvalues=times, title_from_data=True))


def _add_payload_xy_series(
        chart: ScatterChart, data_sheet, last_row: int) -> None:
    """Add the centre-of-mass trajectory as Y values against X values."""
    x_values = Reference(
        data_sheet, min_col=PAYLOAD_CENTER_X_COLUMN,
        min_row=1, max_row=last_row)
    y_values = Reference(
        data_sheet, min_col=PAYLOAD_CENTER_Y_COLUMN,
        min_row=1, max_row=last_row)
    chart.series.append(Series(y_values, xvalues=x_values, title_from_data=True))


@dataclass(frozen=True)
class RotationSample:
    """One synchronized angle and grasp-force sample."""

    elapsed: float
    state: str
    final_target_deg: float
    desired_deg: float
    actual_deg: float
    forces: Dict[str, Tuple[float, float, float]]
    calibrated_forces: Dict[str, Tuple[float, float, float]] = field(
        default_factory=dict)
    sequence: str = ''
    sequence_event: str = ''
    grasp_state: str = ''
    support_removed: Optional[bool] = None
    payload_center_x: Optional[float] = None
    payload_center_y: Optional[float] = None
    payload_roll_deg: Optional[float] = None
    payload_pitch_deg: Optional[float] = None
    payload_yaw_deg: Optional[float] = None


def _force_values(
        sample: RotationSample, point: str
) -> Tuple[Optional[float], ...]:
    force = sample.forces.get(point)
    if force is None:
        return (None, None, None, None, None)
    fx, fy, fz = force
    resultant = math.sqrt(fx * fx + fy * fy + fz * fz)
    return fx, fy, fz, resultant, abs(fx)


def _headers() -> List[str]:
    headers = [
        '時刻 [s]', 'シーケンス', '切替イベント', '把持状態',
        '支持台撤去済み', '回転状態', '最終目標角度 [deg]',
        '指令角度 [deg]', '現在角度 [deg]', '角度誤差 [deg]',
    ]
    for point in FORCE_POINTS:
        label = FORCE_LABELS[point]
        headers.extend([
            f'{label} 生Fx [N]', f'{label} 生Fy [N]', f'{label} 生Fz [N]',
            f'{label} 生合力 [N]', f'{label} 生法線力 |Fx| [N]',
            f'{label} 補正Fx [N]', f'{label} 補正Fy [N]',
            f'{label} 補正Fz [N]', f'{label} 補正合力 [N]',
            f'{label} 補正法線変動 |Fx| [N]',
        ])
    for point in FORCE_POINTS:
        headers.append(f'{FORCE_LABELS[point]} 表示用平滑法線力 [N]')
    headers.extend([
        '物体重心 X [m]', '物体重心 Y [m]',
        '物体姿勢 Roll (X) [deg]', '物体姿勢 Pitch (Y) [deg]',
        '物体姿勢 Yaw (Z) [deg]',
    ])
    return headers


def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    """Load a Japanese-capable font installed with the ROS desktop image."""
    # The regular face keeps the report portable on the target Ubuntu image.
    # Pillow's default font cannot draw Japanese labels.
    del bold
    return ImageFont.truetype(FONT_PATH, size)


def _sequence_segments(samples: List[RotationSample]):
    """Return contiguous (start, end, sequence) segments for the timeline."""
    segments = []
    start = samples[0].elapsed
    sequence = samples[0].sequence or 'ロボット・物体生成'
    for sample in samples[1:]:
        current = sample.sequence or sequence
        if current != sequence:
            segments.append((start, sample.elapsed, sequence))
            start = sample.elapsed
            sequence = current
    # FINISHED is often received only for the final sample.  A zero-width
    # segment cannot hold a readable timeline label, so leave that terminal
    # state to the summary table instead of drawing overlapping text.
    if samples[-1].elapsed > start or not segments:
        segments.append((start, samples[-1].elapsed, sequence))
    return segments


def _draw_dashed_vertical(
        draw: ImageDraw.ImageDraw, x: int, top: int, bottom: int) -> None:
    for y in range(top, bottom, 12):
        draw.line((x, y, x, min(y + 6, bottom)), fill='#C00000', width=2)


def _draw_sequence_force_graph(
        samples: List[RotationSample], path: Path) -> None:
    """Render a self-contained sequence/force chart for embedding in Excel."""
    width, height = 1400, 700
    left, right, top, bottom = 110, 150, 85, 470
    timeline_top, timeline_bottom = 550, 600
    image = Image.new('RGB', (width, height), 'white')
    draw = ImageDraw.Draw(image)
    title_font = _font(25)
    label_font = _font(18)
    small_font = _font(15)
    draw.text((width // 2, 20), '把持力とシーケンス推移',
              font=title_font, fill='#222222', anchor='ma')

    duration = max(5.0, samples[-1].elapsed)
    plot_width = width - left - right

    def x_at(elapsed: float) -> int:
        return left + int(plot_width * elapsed / duration)

    raw_maximum = max(
        (abs(force[0]) for sample in samples for force in sample.forces.values()),
        default=20.0)
    maximum_force = max(25.0, math.ceil(raw_maximum / 5.0) * 5.0)

    def y_at(force: float) -> int:
        return bottom - int(
            (bottom - top) * max(0.0, min(maximum_force, force))
            / maximum_force)

    segments = _sequence_segments(samples)
    for start, end, sequence in segments:
        x_start, x_end = x_at(start), max(x_at(end), x_at(start) + 1)
        color = SEQUENCE_COLORS.get(sequence, '#E7E6E6')
        draw.rectangle((x_start, top, x_end, bottom), fill=color)
        draw.rectangle((x_start, timeline_top, x_end, timeline_bottom),
                       fill=color, outline='#666666', width=1)
        label = sequence.replace('・', '・\n') if x_end - x_start < 130 else sequence
        draw.multiline_text(((x_start + x_end) // 2, (timeline_top + timeline_bottom) // 2),
                            label, font=small_font, fill='#222222',
                            anchor='mm', align='center', spacing=1)

    for value in range(0, int(maximum_force) + 1, 5):
        y = y_at(float(value))
        draw.line((left, y, width - right, y), fill='#C8C8C8', width=1)
        draw.text((left - 12, y), str(value), font=small_font,
                  fill='#333333', anchor='rm')
    draw.line((left, top, left, bottom), fill='#333333', width=2)
    draw.line((left, bottom, width - right, bottom), fill='#333333', width=2)
    draw.text((32, (top + bottom) // 2), '法線力 [N]', font=label_font,
              fill='#222222', anchor='mm')

    for force, color, text in ((17.0, '#548235', '目標 17 N'),
                               (8.0, '#C00000', '下限 8 N')):
        y = y_at(force)
        for x in range(left, width - right, 12):
            draw.line((x, y, min(x + 6, width - right), y), fill=color, width=2)
        if color != '#C00000':
            draw.text((width - right + 10, y), text, font=small_font,
                      fill=color, anchor='lm')

    smoothed = {point: None for point in FORCE_POINTS}
    previous = {point: None for point in FORCE_POINTS}
    for sample in samples:
        x = x_at(sample.elapsed)
        for index, point in enumerate(FORCE_POINTS):
            force = sample.forces.get(point)
            if force is None:
                previous[point] = None
                continue
            normal = abs(force[0])
            before = smoothed[point]
            smoothed[point] = (
                normal if before is None else
                DISPLAY_FORCE_FILTER_ALPHA * normal
                + (1.0 - DISPLAY_FORCE_FILTER_ALPHA) * before)
            y = y_at(smoothed[point])
            if previous[point] is not None:
                draw.line((previous[point][0], previous[point][1], x, y),
                          fill=PLOT_COLORS[index], width=2)
            previous[point] = (x, y)

    legend_x = left + 8
    for index, point in enumerate(FORCE_POINTS):
        y = top + 12 + index * 22
        draw.line((legend_x, y, legend_x + 24, y), fill=PLOT_COLORS[index], width=3)
        draw.text((legend_x + 31, y), FORCE_LABELS[point], font=small_font,
                  fill='#222222', anchor='lm')

    for start, _, sequence in segments[1:]:
        x = x_at(start)
        _draw_dashed_vertical(draw, x, top, bottom)

    draw.text((left, 525), 'シーケンス', font=label_font, fill='#222222')
    for value in range(0, int(duration) + 1, 5):
        x = x_at(float(value))
        draw.text((x, 620), str(value), font=small_font,
                  fill='#333333', anchor='ma')
    draw.text((left + plot_width // 2, 670), 'シミュレーション開始からの時刻 [s]',
              font=label_font, fill='#222222', anchor='ma')
    image.save(path, format='PNG')


def _draw_right_finger_sequence_graph(
        samples: List[RotationSample], path: Path) -> None:
    """Render a sequence chart containing only both AMIR right fingers."""
    width, height = 1400, 700
    left, right, top, bottom = 110, 70, 85, 470
    timeline_top, timeline_bottom = 550, 600
    image = Image.new('RGB', (width, height), 'white')
    draw = ImageDraw.Draw(image)
    title_font = _font(25)
    label_font = _font(18)
    small_font = _font(15)
    draw.text((width // 2, 20), 'AMIR右指の把持力とシーケンス推移',
              font=title_font, fill='#222222', anchor='ma')

    duration = max(5.0, samples[-1].elapsed)
    plot_width = width - left - right

    def x_at(elapsed: float) -> int:
        return left + int(plot_width * elapsed / duration)

    raw_maximum = max(
        (abs(sample.forces[point][0])
         for sample in samples for point in RIGHT_FINGER_POINTS
         if point in sample.forces),
        default=20.0)
    maximum_force = max(25.0, math.ceil(raw_maximum / 5.0) * 5.0)

    def y_at(force: float) -> int:
        return bottom - int(
            (bottom - top) * max(0.0, min(maximum_force, force))
            / maximum_force)

    segments = _sequence_segments(samples)
    for start, end, sequence in segments:
        x_start, x_end = x_at(start), max(x_at(end), x_at(start) + 1)
        color = SEQUENCE_COLORS.get(sequence, '#E7E6E6')
        draw.rectangle((x_start, top, x_end, bottom), fill=color)
        draw.rectangle((x_start, timeline_top, x_end, timeline_bottom),
                       fill=color, outline='#666666', width=1)
        label = sequence.replace('・', '・\n') if x_end - x_start < 130 else sequence
        draw.multiline_text(
            ((x_start + x_end) // 2, (timeline_top + timeline_bottom) // 2),
            label, font=small_font, fill='#222222', anchor='mm',
            align='center', spacing=1)

    for value in range(0, int(maximum_force) + 1, 5):
        y = y_at(float(value))
        draw.line((left, y, width - right, y), fill='#C8C8C8', width=1)
        draw.text((left - 12, y), str(value), font=small_font,
                  fill='#333333', anchor='rm')
    draw.line((left, top, left, bottom), fill='#333333', width=2)
    draw.line((left, bottom, width - right, bottom), fill='#333333', width=2)
    draw.text((32, (top + bottom) // 2), '法線力 [N]', font=label_font,
              fill='#222222', anchor='mm')

    smoothed = {point: None for point in RIGHT_FINGER_POINTS}
    previous = {point: None for point in RIGHT_FINGER_POINTS}
    for sample in samples:
        x = x_at(sample.elapsed)
        for point in RIGHT_FINGER_POINTS:
            force = sample.forces.get(point)
            if force is None:
                previous[point] = None
                continue
            normal = abs(force[0])
            before = smoothed[point]
            smoothed[point] = (
                normal if before is None else
                DISPLAY_FORCE_FILTER_ALPHA * normal
                + (1.0 - DISPLAY_FORCE_FILTER_ALPHA) * before)
            y = y_at(smoothed[point])
            if previous[point] is not None:
                draw.line((previous[point][0], previous[point][1], x, y),
                          fill=RIGHT_FINGER_COLORS[point], width=2)
            previous[point] = (x, y)

    legend_x = left + 8
    for index, point in enumerate(RIGHT_FINGER_POINTS):
        y = top + 12 + index * 22
        draw.line((legend_x, y, legend_x + 24, y),
                  fill=RIGHT_FINGER_COLORS[point], width=3)
        draw.text((legend_x + 31, y), FORCE_LABELS[point], font=small_font,
                  fill='#222222', anchor='lm')

    for start, _, sequence in segments[1:]:
        x = x_at(start)
        _draw_dashed_vertical(draw, x, top, bottom)

    draw.text((left, 525), 'シーケンス', font=label_font, fill='#222222')
    for value in range(
            0, int(duration) + 1, RIGHT_FINGER_TIME_TICK_SECONDS):
        x = x_at(float(value))
        draw.text((x, 625), str(value), font=small_font,
                  fill='#333333', anchor='ma')
    draw.text((left + plot_width // 2, 670),
              'シミュレーション開始からの時刻 [s]',
              font=label_font, fill='#222222', anchor='ma')
    image.save(path, format='PNG')


def write_rotation_workbook(
        path: Path, samples: List[RotationSample],
        generated_at: Optional[datetime] = None) -> Path:
    """Write measurements, summary statistics, and charts to an XLSX file."""
    if not samples:
        raise ValueError('At least one rotation sample is required.')
    output = Path(path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    summary = workbook.active
    summary.title = '概要'
    data_sheet = workbook.create_sheet('時系列データ')
    right_finger_sheet = workbook.create_sheet('右指グラフ')

    generated = generated_at or datetime.now()
    final = samples[-1]
    tracking_errors = [
        abs(sample.desired_deg - sample.actual_deg) for sample in samples]
    summary_rows = [
        ('項目', '値'),
        ('生成日時', generated.strftime('%Y-%m-%d %H:%M:%S')),
        ('記録点数', len(samples)),
        ('記録時間 [s]', final.elapsed),
        ('最終目標角度 [deg]', final.final_target_deg),
        ('最終指令角度 [deg]', final.desired_deg),
        ('最終現在角度 [deg]', final.actual_deg),
        ('最終角度誤差 [deg]', final.desired_deg - final.actual_deg),
        ('最大追従誤差 [deg]', max(tracking_errors)),
        ('最終状態', final.state),
    ]
    for row in summary_rows:
        summary.append(row)

    summary.append([])
    summary.append([
        '把持点（生値）', '最小法線力 [N]',
        '平均法線力 [N]', '最大法線力 [N]'])
    for point in FORCE_POINTS:
        normal_forces = [
            abs(sample.forces[point][0]) for sample in samples
            if point in sample.forces
        ]
        summary.append([
            FORCE_LABELS[point],
            min(normal_forces) if normal_forces else None,
            fmean(normal_forces) if normal_forces else None,
            max(normal_forces) if normal_forces else None,
        ])
    summary.append([])
    summary.append([
        '把持点（把持直後=0補正）', '最小変動 [N]',
        '平均変動 [N]', '最大変動 [N]'])
    for point in FORCE_POINTS:
        changes = [
            abs(sample.calibrated_forces[point][0]) for sample in samples
            if point in sample.calibrated_forces
        ]
        summary.append([
            FORCE_LABELS[point],
            min(changes) if changes else None,
            fmean(changes) if changes else None,
            max(changes) if changes else None,
        ])
    positions = [
        (sample.payload_center_x, sample.payload_center_y)
        for sample in samples
        if (sample.payload_center_x is not None
            and sample.payload_center_y is not None)
    ]
    if positions:
        first_x, first_y = positions[0]
        last_x, last_y = positions[-1]
        summary.append([])
        summary.append(['物体重心位置', '初期値 [m]', '最終値 [m]', '変位 [m]'])
        summary.append(['X', first_x, last_x, last_x - first_x])
        summary.append(['Y', first_y, last_y, last_y - first_y])

    data_sheet.append(_headers())
    smoothed_normal_forces: Dict[str, float] = {}
    for sample in samples:
        row = [
            sample.elapsed,
            sample.sequence,
            sample.sequence_event,
            sample.grasp_state,
            ('はい' if sample.support_removed else 'いいえ'
             if sample.support_removed is not None else ''),
            sample.state,
            sample.final_target_deg,
            sample.desired_deg,
            sample.actual_deg,
            sample.desired_deg - sample.actual_deg,
        ]
        for point in FORCE_POINTS:
            row.extend(_force_values(sample, point))
            calibrated = sample.calibrated_forces.get(point)
            if calibrated is None:
                row.extend((None, None, None, None, None))
            else:
                fx, fy, fz = calibrated
                row.extend((
                    fx, fy, fz,
                    math.sqrt(fx * fx + fy * fy + fz * fz), abs(fx)))
        for point in FORCE_POINTS:
            force = sample.forces.get(point)
            if force is None:
                row.append(None)
                continue
            normal_force = abs(force[0])
            previous = smoothed_normal_forces.get(point, normal_force)
            smoothed = (
                DISPLAY_FORCE_FILTER_ALPHA * normal_force
                + (1.0 - DISPLAY_FORCE_FILTER_ALPHA) * previous)
            smoothed_normal_forces[point] = smoothed
            row.append(smoothed)
        row.extend((
            sample.payload_center_x,
            sample.payload_center_y,
            sample.payload_roll_deg,
            sample.payload_pitch_deg,
            sample.payload_yaw_deg,
        ))
        data_sheet.append(row)

    header_fill = PatternFill('solid', fgColor='1F4E78')
    header_font = Font(color='FFFFFF', bold=True)
    for sheet in (summary, data_sheet):
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
    for force_header_row in (12, 18):
        for cell in summary[force_header_row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
    if positions:
        for cell in summary[24]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

    summary.column_dimensions['A'].width = 25
    summary.column_dimensions['B'].width = 24
    summary.column_dimensions['C'].width = 20
    summary.column_dimensions['D'].width = 20
    data_sheet.freeze_panes = 'A2'
    data_sheet.auto_filter.ref = data_sheet.dimensions
    data_sheet.column_dimensions['A'].width = 13
    data_sheet.column_dimensions['B'].width = 16
    for column in range(3, len(_headers()) + 1):
        data_sheet.column_dimensions[get_column_letter(column)].width = 21
    # Do not assign a style to every numeric cell.  With long recordings this
    # creates a very large worksheet XML and makes each XLSX save expensive.
    # Excel's General format preserves the numeric values and remains usable
    # for filtering, formulas, and charts.

    last_row = data_sheet.max_row
    if last_row >= 2:
        angle_chart = ScatterChart()
        angle_chart.title = '目標角度と現在角度の推移'
        angle_chart.x_axis.title = '時刻 [s]'
        angle_chart.y_axis.title = '角度 [deg]'
        angle_chart.style = 13
        angle_chart.height = 11
        angle_chart.width = 22
        _use_five_unit_axis_labels(angle_chart, final.elapsed)
        angle_times = Reference(data_sheet, min_col=1, min_row=2, max_row=last_row)
        _add_xy_series(angle_chart, data_sheet, (8, 9), angle_times, last_row)
        summary.add_chart(angle_chart, 'F2')

        force_chart = ScatterChart()
        force_chart.title = '把持点の法線力推移（表示用平滑値）'
        force_chart.x_axis.title = '時刻 [s]'
        force_chart.y_axis.title = '法線力 |Fx| [N]'
        force_chart.style = 12
        force_chart.height = 11
        force_chart.width = 22
        _use_five_unit_axis_labels(force_chart, final.elapsed)
        _add_xy_series(
            force_chart, data_sheet, SMOOTHED_NORMAL_FORCE_COLUMNS,
            angle_times, last_row)
        summary.add_chart(force_chart, 'F24')

        calibrated_chart = ScatterChart()
        calibrated_chart.title = '把持直後基準の法線力変動'
        calibrated_chart.x_axis.title = '時刻 [s]'
        calibrated_chart.y_axis.title = '補正法線変動 |Fx| [N]'
        calibrated_chart.style = 10
        calibrated_chart.height = 11
        calibrated_chart.width = 22
        _use_five_unit_axis_labels(calibrated_chart, final.elapsed)
        _add_xy_series(
            calibrated_chart, data_sheet, CALIBRATED_NORMAL_FORCE_COLUMNS,
            angle_times, last_row)
        summary.add_chart(calibrated_chart, 'F46')

        payload_attitude_chart = ScatterChart()
        payload_attitude_chart.title = '物体姿勢 XYZ（Roll・Pitch・Yaw）'
        payload_attitude_chart.x_axis.title = '時刻 [s]'
        payload_attitude_chart.y_axis.title = '姿勢 [deg]'
        payload_attitude_chart.style = 13
        payload_attitude_chart.height = 11
        payload_attitude_chart.width = 22
        _use_five_unit_axis_labels(payload_attitude_chart, final.elapsed)
        _add_xy_series(
            payload_attitude_chart, data_sheet,
            (PAYLOAD_ROLL_COLUMN, PAYLOAD_PITCH_COLUMN, PAYLOAD_YAW_COLUMN),
            angle_times, last_row)
        summary.add_chart(payload_attitude_chart, 'F105')

        payload_position_chart = ScatterChart()
        payload_position_chart.title = '物体重心位置のX-Y軌跡（Zは除外）'
        payload_position_chart.x_axis.title = '物体重心 X [m]'
        payload_position_chart.y_axis.title = '物体重心 Y [m]'
        payload_position_chart.style = 13
        payload_position_chart.height = 11
        payload_position_chart.width = 22
        _add_payload_xy_series(payload_position_chart, data_sheet, last_row)
        summary.add_chart(payload_position_chart, 'F127')

    preview_path: Optional[Path] = None
    right_finger_preview_path: Optional[Path] = None
    try:
        with NamedTemporaryFile(
                prefix='cooperative_rotation_sequence_', suffix='.png',
                delete=False) as temporary:
            preview_path = Path(temporary.name)
        _draw_sequence_force_graph(samples, preview_path)
        preview = ExcelImage(str(preview_path))
        preview.width = 950
        preview.height = 475
        summary.add_image(preview, 'F68')

        with NamedTemporaryFile(
                prefix='cooperative_rotation_right_fingers_', suffix='.png',
                delete=False) as temporary:
            right_finger_preview_path = Path(temporary.name)
        _draw_right_finger_sequence_graph(samples, right_finger_preview_path)
        right_finger_preview = ExcelImage(str(right_finger_preview_path))
        right_finger_preview.width = 1120
        right_finger_preview.height = 560
        right_finger_sheet.add_image(right_finger_preview, 'A1')
        right_finger_sheet.sheet_view.showGridLines = False
        workbook.save(output)
    finally:
        if preview_path is not None:
            preview_path.unlink(missing_ok=True)
        if right_finger_preview_path is not None:
            right_finger_preview_path.unlink(missing_ok=True)
    return output
