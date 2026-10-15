"""Record distributed straight transport in the grasp-force Excel format."""

from datetime import datetime
import math
from pathlib import Path

import rclpy
from geometry_msgs.msg import PoseStamped, Twist, WrenchStamped
from nav_msgs.msg import Odometry
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rclpy.node import Node
from rclpy.qos import DurabilityPolicy, QoSProfile
from std_msgs.msg import Bool, String
from std_srvs.srv import Trigger

from .excel_report import RotationSample, write_rotation_workbook
from .transform_utils import quaternion_to_rpy


DETAIL_HEADERS = [
    '時刻 [s]', 'リーダー状態', 'フォロア状態', '安全状態', '滑り状態',
    '目標X [m]', '目標Y [m]', '物体X [m]', '物体Y [m]', '目標位置誤差 [m]',
    'AMIR1 X [m]', 'AMIR1 Y [m]', 'AMIR1 Yaw [deg]',
    'AMIR2 X [m]', 'AMIR2 Y [m]', 'AMIR2 Yaw [deg]',
    'AMIR1 指令Vx [m/s]', 'AMIR1 指令Vy [m/s]',
    'AMIR2 指令Vx [m/s]', 'AMIR2 指令Vy [m/s]',
]


class DistributedDataRecorder(Node):
    """Sample the two robots and payload on one simulation-time axis."""

    def __init__(self):
        super().__init__('distributed_transport_data_recorder')
        self.declare_parameter('goal_x', 0.5)
        self.declare_parameter('goal_y', 0.0)
        self.declare_parameter('target_normal_force', 20.0)
        self.declare_parameter('sample_rate', 10.0)
        self.declare_parameter(
            'output_directory', '~/ros2_humble_ws/rotation_measurements')
        self.declare_parameter('output_file', '')
        self.goal = (
            float(self.get_parameter('goal_x').value),
            float(self.get_parameter('goal_y').value))
        rate = float(self.get_parameter('sample_rate').value)
        if not math.isfinite(rate) or rate <= 0.0:
            raise ValueError('sample_rate must be positive and finite')

        self.leader_state = 'WAITING'
        self.follower_state = 'WAITING'
        self.grasp_state = 'WAITING_FOR_STATUS'
        self.safety_status = ''
        self.slip_status = ''
        self.support_removed = None
        self.payload = None
        self.payload_rpy = None
        self.odom = {}
        self.commands = {}
        self.forces = {}
        self.calibrated_forces = {}
        self.samples = []
        self.details = []
        self.started_at = None
        self.last_sequence = None
        self.has_moved = False
        self.complete = False
        self.saved_path = None
        self.saved_count = 0

        self.create_subscription(
            String, '/distributed_transport/leader/state',
            self._on_leader_state, 20)
        self.create_subscription(
            String, '/distributed_transport/follower/state',
            lambda msg: setattr(self, 'follower_state', msg.data), 20)
        self.create_subscription(
            String, '/cooperative_transport/friction_grasp_status',
            lambda msg: setattr(
                self, 'grasp_state', msg.data.split(';', 1)[0].strip()), 20)
        self.create_subscription(
            String, '/cooperative_transport/safety_status',
            lambda msg: setattr(self, 'safety_status', msg.data), 20)
        self.create_subscription(
            String, '/cooperative_transport/slip_status',
            lambda msg: setattr(self, 'slip_status', msg.data), 20)
        self.create_subscription(
            PoseStamped, '/cooperative_transport/measured_payload_pose',
            self._on_payload, 20)
        latched = QoSProfile(depth=1)
        latched.durability = DurabilityPolicy.TRANSIENT_LOCAL
        self.create_subscription(
            Bool, '/cooperative_transport/support_removed',
            lambda msg: setattr(self, 'support_removed', msg.data), latched)
        for robot in ('amir1', 'amir2'):
            self.create_subscription(
                Odometry, f'/{robot}/odom',
                lambda msg, key=robot: self._on_odom(key, msg), 20)
            self.create_subscription(
                Twist, f'/{robot}/rover_twist',
                lambda msg, key=robot: self._on_command(key, msg), 20)
            for side in ('left', 'right'):
                point = f'{robot}_{side}'
                self.create_subscription(
                    WrenchStamped, f'/{robot}/finger_{side}_wrench',
                    lambda msg, key=point: self._on_force(key, msg), 20)
                self.create_subscription(
                    WrenchStamped,
                    f'/{robot}/finger_{side}_wrench_calibrated',
                    lambda msg, key=point: self._on_force(key, msg, True), 20)
        self.report_pub = self.create_publisher(
            String, '/distributed_transport/excel_report', latched)
        self.create_service(
            Trigger, '/distributed_transport/save_excel', self._on_save)
        self.create_timer(1.0 / rate, self._sample)
        self.get_logger().info(
            'Distributed transport Excel recorder ready; auto-saves on finish.')

    def _now(self):
        return self.get_clock().now().nanoseconds * 1.0e-9

    def _on_leader_state(self, msg):
        self.leader_state = msg.data
        if msg.data in ('MOVING', 'TRACKING_PAUSED'):
            self.has_moved = True
        if (msg.data == 'FINISHED' or
                (msg.data == 'WAITING_OR_STOPPED' and self.has_moved)):
            if not self.complete:
                self._sample()
                self.complete = True
                self.save()

    def _on_payload(self, msg):
        pose = msg.pose
        self.payload = (pose.position.x, pose.position.y)
        q = pose.orientation
        self.payload_rpy = tuple(math.degrees(value) for value in
                                 quaternion_to_rpy((q.x, q.y, q.z, q.w)))

    def _on_odom(self, robot, msg):
        pose = msg.pose.pose
        q = pose.orientation
        yaw = quaternion_to_rpy((q.x, q.y, q.z, q.w))[2]
        self.odom[robot] = (
            pose.position.x, pose.position.y, math.degrees(yaw))

    def _on_command(self, robot, msg):
        self.commands[robot] = (msg.linear.x, msg.linear.y)

    def _on_force(self, point, msg, calibrated=False):
        force = msg.wrench.force
        destination = self.calibrated_forces if calibrated else self.forces
        destination[point] = (force.x, force.y, force.z)

    def _sequence(self):
        if self.leader_state == 'FINISHED':
            return '搬送完了'
        if self.leader_state == 'WAITING_OR_STOPPED' and self.has_moved:
            return '異常終了'
        if self.leader_state in ('MOVING', 'TRACKING_PAUSED'):
            return '搬送（直進）'
        if self.support_removed:
            return '把持完了・搬送準備'
        if self.grasp_state in ('CLOSING_GRIPPERS', 'LIFTING', 'HOLDING'):
            return '把持'
        return 'ロボット・物体生成'

    def _sample(self):
        if self.complete:
            return
        now = self._now()
        if self.started_at is None:
            self.started_at = now
        elapsed = max(0.0, now - self.started_at)
        sequence = self._sequence()
        event = (f'開始: {sequence}' if sequence != self.last_sequence else '')
        self.last_sequence = sequence
        x, y = self.payload if self.payload is not None else (None, None)
        roll, pitch, yaw = (
            self.payload_rpy if self.payload_rpy is not None else
            (None, None, None))
        self.samples.append(RotationSample(
            elapsed=elapsed, state=self.leader_state,
            final_target_deg=0.0, desired_deg=0.0,
            actual_deg=yaw if yaw is not None else 0.0,
            forces=dict(self.forces),
            calibrated_forces=dict(self.calibrated_forces),
            sequence=sequence, sequence_event=event,
            grasp_state=self.grasp_state,
            support_removed=self.support_removed,
            payload_center_x=x, payload_center_y=y,
            payload_roll_deg=roll, payload_pitch_deg=pitch,
            payload_yaw_deg=yaw))
        error = (math.hypot(self.goal[0] - x, self.goal[1] - y)
                 if x is not None and y is not None else None)
        row = [elapsed, self.leader_state, self.follower_state,
               self.safety_status, self.slip_status,
               self.goal[0], self.goal[1], x, y, error]
        for robot in ('amir1', 'amir2'):
            row.extend(self.odom.get(robot, (None, None, None)))
        for robot in ('amir1', 'amir2'):
            row.extend(self.commands.get(robot, (None, None)))
        self.details.append(row)

    def _decorate_workbook(self, workbook):
        summary = workbook['概要']
        series = workbook['時系列データ']
        for column, label in {
                6: '直進状態', 7: '目標Yaw [deg]',
                8: '指令Yaw [deg]', 9: '物体Yaw [deg]',
                10: 'Yaw誤差 [deg]'}.items():
            series.cell(1, column).value = label
        for row, label in {
                5: '目標Yaw [deg]', 6: '指令Yaw [deg]',
                7: '物体Yaw [deg]', 8: 'Yaw誤差 [deg]',
                9: '最大Yaw誤差 [deg]'}.items():
            summary.cell(row, 1).value = label
        if summary._charts:
            summary._charts[0].title = '目標Yawと物体Yawの推移'
        summary.append([])
        summary.append(['分散型直進搬送', '値'])
        summary.append(['目標X [m]', self.goal[0]])
        summary.append(['目標Y [m]', self.goal[1]])
        final_error = self.details[-1][9]
        summary.append(['最終位置誤差 [m]', final_error])
        summary.append(['リーダー最終状態', self.leader_state])
        summary.append(['フォロア最終状態', self.follower_state])
        summary.append(['安全最終状態', self.safety_status])
        summary.append(['滑り最終状態', self.slip_status])

        sheet = workbook.create_sheet('分散搬送データ')
        sheet.append(DETAIL_HEADERS)
        for row in self.details:
            sheet.append(row)
        header_fill = PatternFill('solid', fgColor='1F4E78')
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center')
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions['A'].width = 13
        for col in range(2, len(DETAIL_HEADERS) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 23
        if sheet.max_row >= 2:
            chart = ScatterChart()
            chart.title = '物体位置と目標位置の推移'
            chart.x_axis.title = '時刻 [s]'
            chart.y_axis.title = 'X位置 [m]'
            chart.height = 11
            chart.width = 22
            times = Reference(sheet, min_col=1, min_row=2,
                              max_row=sheet.max_row)
            for col in (6, 8):
                values = Reference(sheet, min_col=col, min_row=1,
                                   max_row=sheet.max_row)
                chart.series.append(Series(
                    values, xvalues=times, title_from_data=True))
            sheet.add_chart(chart, 'V2')

    def _output_path(self):
        configured = str(self.get_parameter('output_file').value).strip()
        if configured:
            return Path(configured).expanduser().resolve()
        directory = Path(str(
            self.get_parameter('output_directory').value)).expanduser()
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        return (directory / f'distributed_straight_{timestamp}.xlsx').resolve()

    def save(self):
        if not self.samples:
            return None
        if self.saved_path is None:
            self.saved_path = self._output_path()
        if (self.saved_path.exists() and
                self.saved_count == len(self.samples)):
            return self.saved_path
        write_rotation_workbook(
            self.saved_path, self.samples,
            target_normal_force=float(
                self.get_parameter('target_normal_force').value),
            customize=self._decorate_workbook)
        self.saved_count = len(self.samples)
        self.get_logger().info(f'Excel report saved: {self.saved_path}')
        if rclpy.ok():
            self.report_pub.publish(String(data=str(self.saved_path)))
        return self.saved_path

    def _on_save(self, _request, response):
        path = self.save()
        response.success = path is not None
        response.message = str(path) if path else 'No samples recorded yet.'
        return response


def main(args=None):
    rclpy.init(args=args)
    node = DistributedDataRecorder()
    try:
        rclpy.spin(node)
    except KeyboardInterrupt:
        pass
    finally:
        node.save()
        node.destroy_node()
        if rclpy.ok():
            rclpy.shutdown()


if __name__ == '__main__':
    main()
