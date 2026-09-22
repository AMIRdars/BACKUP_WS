import math

from openpyxl import load_workbook

from cooperative_rotation.excel_report import (
    FORCE_POINTS,
    RotationSample,
    write_rotation_workbook,
)


def test_workbook_contains_data_summary_and_charts(tmp_path):
    forces = {
        point: (100.0 + index, 2.0, -3.0)
        for index, point in enumerate(FORCE_POINTS)
    }
    calibrated = {
        point: (1.0 + index, 0.2, -0.3)
        for index, point in enumerate(FORCE_POINTS)
    }
    samples = [
        RotationSample(
            0.0, 'ROTATING', 90.0, 0.0, 0.0, forces, calibrated,
            '把持', '開始: 把持', 'HOLDING', False,
            0.10, -0.20, 0.0, 0.0, 0.0),
        RotationSample(
            1.0, 'FINISHED', 90.0, 90.0, 89.95, forces, calibrated,
            '搬送完了', '開始: 搬送完了', 'HOLDING', True,
            0.20, -0.10, 1.0, -2.0, 89.95),
    ]
    output = write_rotation_workbook(tmp_path / 'rotation.xlsx', samples)

    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ['概要', '時系列データ', '右指グラフ']
    assert workbook['時系列データ'].max_row == 3
    assert workbook['時系列データ']['I3'].value == 89.95
    assert workbook['概要']['B7'].value == 89.95
    assert workbook['時系列データ']['T3'].value == 1.0
    assert workbook['時系列データ'].cell(2, 51).value == 100.0
    assert workbook['時系列データ'].cell(2, 54).value == 103.0
    assert workbook['時系列データ']['B2'].value == '把持'
    assert workbook['時系列データ']['C3'].value == '開始: 搬送完了'
    assert workbook['時系列データ'].cell(3, 55).value == 0.20
    assert workbook['時系列データ'].cell(3, 56).value == -0.10
    assert workbook['時系列データ'].cell(3, 59).value == 89.95
    assert len(workbook['概要']._charts) == 5
    assert len(workbook['概要']._images) == 1
    assert len(workbook['右指グラフ']._images) == 1
    assert math.isclose(workbook['概要']['B8'].value, 0.05)
    for chart in workbook['概要']._charts[:4]:
        assert chart.x_axis.numFmt.formatCode == '0'
        assert chart.y_axis.numFmt.formatCode == '0'
        assert chart.x_axis.majorUnit == 5.0
        assert chart.y_axis.majorUnit == 5.0
